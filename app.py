
import streamlit as st
from google import genai
from google.genai import types
from pypdf import PdfReader
from docx import Document
from pptx import Presentation
from openpyxl import load_workbook
import time
import re
from io import BytesIO
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer

# =========================================================
# PAGE CONFIGURATION
# =========================================================

st.set_page_config(
    page_title="EduAgent AI",
    page_icon="🎓",
    layout="wide"
)


# =========================================================
# GEMINI API
# =========================================================

api_key = st.secrets["GEMINI_API_KEY"]
client = genai.Client(api_key=api_key)


# =========================================================
# SESSION STATE
# =========================================================

defaults = {
    "document_text": "",
    "document_name": "",
    "document_chunks": [],
    "image_files": [],
    "processed_files": [],
    "chat_history": [],
    "question_count": 0,
    "knowledge_loaded": False
}

for key, value in defaults.items():
    if key not in st.session_state:
        st.session_state[key] = value


# =========================================================
# GEMINI RETRY FUNCTION
# =========================================================

def generate_with_retry(contents, max_retries=3):

    for attempt in range(max_retries):

        try:

            response = client.models.generate_content(
                model="gemini-3.5-flash-lite",
                contents=contents
            )

            return response.text

        except Exception as e:

            error_text = str(e)

            if "503" in error_text or "UNAVAILABLE" in error_text:

                if attempt < max_retries - 1:
                    time.sleep(5)
                else:
                    raise Exception(
                        "Gemini is still unavailable after several attempts."
                    )

            else:
                raise e

    raise Exception(
        "Gemini is still unavailable after several attempts."
    )

# =========================================================
# TEXT CLEANING
# =========================================================

def clean_text(text):

    if not text:
        return ""

    return re.sub(r"\s+", " ", text).strip()


# =========================================================
# PDF EXTRACTION
# =========================================================

def extract_pdf_text(uploaded_file):

    reader = PdfReader(uploaded_file)

    text = ""

    for page in reader.pages:

        page_text = page.extract_text()

        if page_text:
            text += page_text + "\n"

    return text


# =========================================================
# DOCX EXTRACTION
# =========================================================

def extract_docx_text(uploaded_file):

    document = Document(uploaded_file)

    text_parts = []

    for paragraph in document.paragraphs:

        if paragraph.text.strip():
            text_parts.append(paragraph.text)

    for table in document.tables:

        for row in table.rows:

            row_text = []

            for cell in row.cells:
                row_text.append(cell.text.strip())

            text_parts.append(" | ".join(row_text))

    return "\n".join(text_parts)


# =========================================================
# TXT EXTRACTION
# =========================================================

def extract_txt_text(uploaded_file):

    raw_data = uploaded_file.read()

    for encoding in ["utf-8", "utf-8-sig", "latin-1"]:

        try:
            return raw_data.decode(encoding)

        except UnicodeDecodeError:
            continue

    return raw_data.decode("utf-8", errors="ignore")


# =========================================================
# PPTX EXTRACTION
# =========================================================

def extract_pptx_text(uploaded_file):

    presentation = Presentation(uploaded_file)

    text_parts = []

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1
    ):

        text_parts.append(
            f"Slide {slide_number}"
        )

        for shape in slide.shapes:

            if hasattr(shape, "text"):

                if shape.text.strip():
                    text_parts.append(shape.text)

    return "\n".join(text_parts)


# =========================================================
# XLSX EXTRACTION
# =========================================================

def extract_xlsx_text(uploaded_file):

    workbook = load_workbook(
        uploaded_file,
        read_only=True,
        data_only=True
    )

    text_parts = []

    for sheet in workbook.worksheets:

        text_parts.append(
            f"Sheet: {sheet.title}"
        )

        for row in sheet.iter_rows(values_only=True):

            values = []

            for value in row:

                if value is not None:
                    values.append(str(value))

            if values:
                text_parts.append(
                    " | ".join(values)
                )

    return "\n".join(text_parts)

# =========================================================
# TEXT CHUNKING
# =========================================================

def create_chunks(
    text,
    chunk_size=1800,
    overlap=300
):

    text = clean_text(text)

    if not text:
        return []

    chunks = []

    start = 0

    while start < len(text):

        end = start + chunk_size

        chunk = text[start:end]

        if chunk.strip():
            chunks.append(chunk.strip())

        if end >= len(text):
            break

        start = end - overlap

    return chunks


# =========================================================
# SIMPLE RAG RETRIEVAL
# =========================================================

def retrieve_relevant_chunks(
    question,
    chunks,
    top_k=4
):

    if not chunks:
        return []

    question_words = set(
        re.findall(
            r"\b[a-zA-Z0-9]+\b",
            question.lower()
        )
    )

    scored_chunks = []

    for chunk in chunks:

        chunk_words = set(
            re.findall(
                r"\b[a-zA-Z0-9]+\b",
                chunk.lower()
            )
        )

        overlap = len(
            question_words.intersection(chunk_words)
        )

        score = overlap

        if question.lower() in chunk.lower():
            score += 10

        scored_chunks.append(
            (score, chunk)
        )

    scored_chunks.sort(
        key=lambda x: x[0],
        reverse=True
    )

    return [
        chunk
        for score, chunk in scored_chunks[:top_k]
        if score > 0
    ]


# =========================================================
# PROCESS MULTIPLE UPLOADED FILES
# =========================================================

def process_uploaded_files(uploaded_files):

    document_text_parts = []
    processed_files = []
    image_files = []

    image_count = 0

    for uploaded_file in uploaded_files:

        filename = uploaded_file.name
        extension = filename.lower().split(".")[-1]

        if extension == "pdf":

            text = extract_pdf_text(uploaded_file)

            if text.strip():

                document_text_parts.append(
                    f"\n===== FILE: {filename} =====\n{text}"
                )

                processed_files.append(filename)

        elif extension == "docx":

            text = extract_docx_text(uploaded_file)

            if text.strip():

                document_text_parts.append(
                    f"\n===== FILE: {filename} =====\n{text}"
                )

                processed_files.append(filename)

        elif extension == "txt":

            text = extract_txt_text(uploaded_file)

            if text.strip():

                document_text_parts.append(
                    f"\n===== FILE: {filename} =====\n{text}"
                )

                processed_files.append(filename)

        elif extension == "pptx":

            text = extract_pptx_text(uploaded_file)

            if text.strip():

                document_text_parts.append(
                    f"\n===== FILE: {filename} =====\n{text}"
                )

                processed_files.append(filename)

        elif extension == "xlsx":

            text = extract_xlsx_text(uploaded_file)

            if text.strip():

                document_text_parts.append(
                    f"\n===== FILE: {filename} =====\n{text}"
                )

                processed_files.append(filename)

        elif extension in ["png", "jpg", "jpeg"]:

            image_count += 1

            if image_count > 10:

                raise ValueError(
                    "Maximum 10 images are allowed."
                )

            image_data = uploaded_file.getvalue()

            if extension == "png":
                mime_type = "image/png"
            else:
                mime_type = "image/jpeg"

            image_files.append(
                {
                    "name": filename,
                    "data": image_data,
                    "mime_type": mime_type
                }
            )

            processed_files.append(filename)

    combined_text = "\n".join(
        document_text_parts
    )

    chunks = create_chunks(
        combined_text
    )

    return (
        combined_text,
        chunks,
        image_files,
        processed_files
    )

# =========================================================
# DOCUMENT QUESTION ANSWERING
# =========================================================

def answer_document_question(question):

    chunks = st.session_state.document_chunks
    image_files = st.session_state.image_files

    relevant_chunks = retrieve_relevant_chunks(
        question,
        chunks,
        top_k=4
    )

    retrieved_text = ""

    if relevant_chunks:

        for index, chunk in enumerate(
            relevant_chunks,
            start=1
        ):

            retrieved_text += (
                f"\n\n--- Source Section {index} ---\n"
                f"{chunk}"
            )

    else:

        retrieved_text = (
            "No relevant text section was retrieved."
        )

    prompt = f"""
You are EduAgent AI, an educational assistant.

The user has uploaded study materials.
These materials may include PDF, DOCX, TXT, PPTX,
XLSX documents and images.

User question:
{question}

Retrieved text from uploaded study materials:
{retrieved_text}

There may also be uploaded images attached after
this instruction.

IMPORTANT SOURCE RULES:

1. First check the uploaded study material.

2. If the answer is clearly supported by the uploaded
material, answer using that material.

3. Do not invent information and pretend it came from
the uploaded material.

4. If the answer is NOT found or cannot be reliably
determined from the uploaded material, clearly begin
your response with:

📄 This information was not found in the uploaded document.

Then provide a useful answer from your general knowledge
and clearly begin that part with:

🤖 Answer from EduAgent's general knowledge:

5. If the uploaded material contains the answer,
begin with:

📄 Answer from uploaded document:

6. If an image contains relevant information, you may
use the image as part of the uploaded study material.

7. Keep answers clear and suitable for students and
teachers.

8. Never claim that general knowledge came from the
uploaded material.

Answer the user's question now.
"""

    contents = [prompt]

    for image in image_files:

        contents.append(
            types.Part.from_bytes(
                data=image["data"],
                mime_type=image["mime_type"]
            )
        )

    answer = generate_with_retry(
        contents
    )

    return answer, relevant_chunks

# =========================================================
# PLANNING AGENT
# =========================================================

def planning_agent(
    subject,
    grade,
    topic,
    duration,
    difficulty
):

    prompt = f"""
You are the Planning Agent of EduAgent AI,
a multi-agent teaching assistant for teachers.

Create a structured lesson plan using the information below.

Subject: {subject}
Grade: {grade}
Topic: {topic}
Class Duration: {duration} minutes
Difficulty Level: {difficulty}

Create:

1. Three clear learning objectives
2. A short introduction to the topic
3. A lesson sequence with time allocation
4. Teaching/learning activities
5. A suitable classroom activity
6. A short conclusion
7. Materials/resources required

Make the lesson appropriate for the specified grade.
Use simple, clear language.
Focus on accurate educational content.

Return the answer with clear headings.
"""

    return generate_with_retry(prompt)


# =========================================================
# CONTENT AGENT
# =========================================================

def content_agent(
    subject,
    grade,
    topic,
    lesson_plan
):

    prompt = f"""
You are the Content Agent of EduAgent AI.

Your job is to create high-quality teaching content
based on the lesson plan prepared by the Planning Agent.

Teacher information:
Subject: {subject}
Grade: {grade}
Topic: {topic}

Lesson plan from Planning Agent:
{lesson_plan}

Create:

1. A clear explanation of the topic
2. Important concepts and definitions
3. Important formulas, if applicable
4. One or two simple real-life examples
5. A worked example suitable for the students
6. A classroom activity
7. Three questions the teacher can ask students

Requirements:
- Keep the content appropriate for the specified grade.
- Use simple language.
- For Physics, make sure formulas and scientific concepts are accurate.
- Do not introduce advanced concepts unnecessarily.
- Organize the answer with clear headings.
"""

    return generate_with_retry(prompt)


# =========================================================
# ASSESSMENT AGENT
# =========================================================

def assessment_agent(
    subject,
    grade,
    topic,
    content
):

    prompt = f"""
You are the Assessment Agent of EduAgent AI.

Create assessments based on the teaching content.

Subject: {subject}
Grade: {grade}
Topic: {topic}

Teaching content:
{content}

Create:

1. Five multiple-choice questions (MCQs)
   - Give 4 options for each question
   - Clearly identify the correct answer

2. Three short-answer questions
   - Suitable for the specified grade
   - Include a brief answer/key point for each

3. Two homework questions
   - One basic question
   - One application-based question

Requirements:
- Questions must be based on the provided content.
- Keep the difficulty appropriate for the grade.
- Avoid ambiguous questions.
- Make sure the answers are scientifically accurate.
- Use clear formatting and headings.
"""

    return generate_with_retry(prompt)


# =========================================================
# REVIEW AGENT
# =========================================================

def review_agent(
    subject,
    grade,
    topic,
    lesson_plan,
    content,
    assessment
):

    prompt = f"""
You are the Review Agent of EduAgent AI.

Review the complete teaching package.

Subject: {subject}
Grade: {grade}
Topic: {topic}

LESSON PLAN:
{lesson_plan}

TEACHING CONTENT:
{content}

ASSESSMENT:
{assessment}

Review for:

1. Scientific accuracy
2. Grade appropriateness
3. Alignment between learning objectives and content
4. Alignment between content and assessment
5. Clarity and readability
6. Correct formulas and calculations
7. Quality of MCQs and answers
8. Suitable difficulty level

Provide:

A. Overall quality rating out of 10

B. Problems or issues found

C. Specific corrections needed

D. Final recommendations for the teacher

If everything is correct, clearly say:
"No major issues found."

Keep the review concise and useful.
"""

    return generate_with_retry(prompt)


# =========================================================
# COMPLETE MULTI-AGENT PIPELINE
# =========================================================

def generate_teaching_package(
    subject,
    grade,
    topic,
    duration,
    difficulty
):

    lesson_plan = planning_agent(
        subject,
        grade,
        topic,
        duration,
        difficulty
    )

    content = content_agent(
        subject,
        grade,
        topic,
        lesson_plan
    )

    assessment = assessment_agent(
        subject,
        grade,
        topic,
        content
    )

    review = review_agent(
        subject,
        grade,
        topic,
        lesson_plan,
        content,
        assessment
    )

    return {
        "lesson_plan": lesson_plan,
        "content": content,
        "assessment": assessment,
        "review": review
    }

# =========================================================
# SIMPLE STREAMLIT UI
# =========================================================

st.title("🎓 EduAgent AI")

st.write(
    "Multi-Agent Teaching Assistant — create a complete "
    "classroom-ready teaching package and interact with "
    "your uploaded study materials."
)

st.markdown("### 🤖 How EduAgent AI works")

agent_cols = st.columns(4)

agent_info = [
    ("📋", "Planning Agent", "Builds the lesson structure"),
    ("🧠", "Content Agent", "Creates teaching material"),
    ("📝", "Assessment Agent", "Creates questions & homework"),
    ("🔍", "Review Agent", "Checks quality and alignment")
]

for col, (icon, name, description) in zip(agent_cols, agent_info):
    with col:
        st.markdown(f"### {icon}")
        st.write(f"**{name}**")
        st.caption(description)


# =========================================================
# MAIN TABS
# =========================================================

tab_knowledge, tab_teaching = st.tabs(
    [
        "💬 Ask Your Knowledge Base",
        "🧑‍🏫 Teaching Package"
    ]
)


# =========================================================
# KNOWLEDGE BASE TAB
# =========================================================

with tab_knowledge:

    st.markdown(
        "### 📚 Upload Your Study Materials"
    )

    st.write(
        "Upload multiple documents and images. "
        "EduAgent will use the uploaded material "
        "to answer questions."
    )

    uploaded_files = st.file_uploader(
        "Upload your study materials",
        type=[
            "pdf",
            "docx",
            "txt",
            "pptx",
            "xlsx",
            "png",
            "jpg",
            "jpeg"
        ],
        accept_multiple_files=True,
        help=(
            "Supported: PDF, DOCX, TXT, PPTX, XLSX, "
            "PNG, JPG, JPEG. Maximum 10 images."
        )
    )

    if uploaded_files:

        image_upload_count = sum(
            1
            for file in uploaded_files
            if file.name.lower().split(".")[-1]
            in ["png", "jpg", "jpeg"]
        )

        if image_upload_count > 10:

            st.error(
                "❌ Maximum 10 images are allowed. "
                f"You selected {image_upload_count}."
            )

        else:

            if st.button(
                "📥 Process & Load Knowledge Base",
                type="primary",
                use_container_width=True
            ):

                try:

                    with st.spinner(
                        "📚 Processing your study materials..."
                    ):

                        (
                            combined_text,
                            chunks,
                            image_files,
                            processed_files
                        ) = process_uploaded_files(
                            uploaded_files
                        )

                    st.session_state.document_text = (
                        combined_text
                    )

                    st.session_state.document_chunks = (
                        chunks
                    )

                    st.session_state.image_files = (
                        image_files
                    )

                    st.session_state.processed_files = (
                        processed_files
                    )

                    st.session_state.document_name = (
                        ", ".join(processed_files)
                    )

                    st.session_state.chat_history = []

                    st.session_state.question_count = 0

                    st.session_state.knowledge_loaded = True

                    st.success(
                        "✅ Knowledge base loaded successfully!"
                    )

                except Exception as e:

                    st.error(
                        f"❌ Error while processing files: {e}"
                    )


    # =====================================================
    # KNOWLEDGE STATUS
    # =====================================================

    st.markdown(
        "### 📊 Knowledge Base Status"
    )

    status_col1, status_col2, status_col3, status_col4 = st.columns(4)

    with status_col1:

        if st.session_state.knowledge_loaded:

            st.metric(
                "Status",
                "Loaded"
            )

        else:

            st.metric(
                "Status",
                "Not Loaded"
            )

    with status_col2:

        st.metric(
            "Files",
            len(
                st.session_state.processed_files
            )
        )

    with status_col3:

        st.metric(
            "Text Sections",
            len(
                st.session_state.document_chunks
            )
        )

    with status_col4:

        st.metric(
            "Images",
            f"{len(st.session_state.image_files)}/10"
        )


    # =====================================================
    # LOADED FILES
    # =====================================================

    if st.session_state.processed_files:

        with st.expander(
            "📁 View Loaded Files"
        ):

            for filename in (
                st.session_state.processed_files
            ):

                st.write(
                    f"✅ {filename}"
                )


    # =====================================================
    # ASK QUESTION
    # =====================================================

    if st.session_state.knowledge_loaded:

        st.markdown(
            "### 💬 Ask Your Uploaded Material"
        )

        question = st.text_input(
            "Ask a question",
            placeholder=(
                "Example: What is Ohm's Law?"
            )
        )

        ask_button = st.button(
            "🔎 Ask EduAgent",
            type="primary"
        )

        if ask_button:

            if not question.strip():

                st.warning(
                    "⚠️ Please enter a question."
                )

            else:

                try:

                    with st.spinner(
                        "🤖 Searching your study materials..."
                    ):

                        (
                            answer,
                            relevant_chunks
                        ) = answer_document_question(
                            question
                        )

                    st.session_state.chat_history.append(
                        {
                            "question": question,
                            "answer": answer,
                            "sources": relevant_chunks
                        }
                    )

                    st.session_state.question_count += 1

                except Exception as e:

                    st.error(
                        f"❌ Error: {e}"
                    )


    # =====================================================
    # CHAT HISTORY
    # =====================================================

    if st.session_state.chat_history:

        st.markdown(
            "### 🗨️ Questions & Answers"
        )

        for chat in reversed(
            st.session_state.chat_history
        ):

            with st.chat_message("user"):

                st.write(
                    chat["question"]
                )

            with st.chat_message("assistant"):

                st.markdown(
                    chat["answer"]
                )

                if chat["sources"]:

                    with st.expander(
                        "📖 View Source Context"
                    ):

                        for index, source in enumerate(
                            chat["sources"],
                            start=1
                        ):

                            st.markdown(
                                f"**Source Section {index}**"
                            )

                            st.write(
                                source
                            )

                            st.divider()


    # =====================================================
    # IMAGE PREVIEW
    # =====================================================

    if st.session_state.image_files:

        with st.expander(
            "🖼️ View Uploaded Images"
        ):

            for image in (
                st.session_state.image_files
            ):

                st.markdown(
                    f"**{image['name']}**"
                )

                st.image(
                    image["data"],
                    use_container_width=True
                )


# =========================================================
# TEACHING PACKAGE TAB
# =========================================================

with tab_teaching:

    st.markdown(
        "### 🧑‍🏫 Create Your Teaching Package"
    )

    st.markdown(
        """
        <div class="section-note">
        Enter your lesson details below, then let the
        four AI agents prepare the package.
        </div>
        """,
        unsafe_allow_html=True
    )

    with st.form(
        "teaching_package_form"
    ):

        input_col1, input_col2 = st.columns(2)

        with input_col1:

            subject = st.text_input(
                "Subject",
                value="Physics",
                key="subject_input",
                placeholder="e.g. Physics"
            )

            grade = st.text_input(
                "Grade",
                value="Grade 9",
                key="grade_input",
                placeholder="e.g. Grade 9"
            )

            topic = st.text_input(
                "Topic",
                value="Ohm's Law",
                key="topic_input",
                placeholder="e.g. Ohm's Law"
            )

        with input_col2:

            duration = st.number_input(
                "Class Duration (minutes)",
                min_value=10,
                max_value=180,
                value=40,
                step=5,
                key="duration_input"
            )

            difficulty = st.selectbox(
                "Difficulty Level",
                [
                    "Easy",
                    "Medium",
                    "Hard"
                ],
                index=1,
                key="difficulty_input"
            )

            st.caption(
                "💡 Choose a difficulty level appropriate "
                "for your students."
            )

        submitted = st.form_submit_button(
            "🚀 Generate Teaching Package",
            type="primary",
            use_container_width=True
        )


    # =====================================================
    # GENERATE PACKAGE
    # =====================================================

    if submitted:

        with st.spinner(
            "🤖 Four AI agents are preparing "
            "your teaching package..."
        ):

            try:

                package = generate_teaching_package(
                    subject,
                    grade,
                    topic,
                    duration,
                    difficulty
                )

            except Exception as e:

                st.error(
                    f"❌ Error while generating package: {e}"
                )

                st.stop()

        st.success(
            "✅ Teaching package generated successfully!"
        )

        st.markdown(
            f"### 📚 Teaching Package: {topic}"
        )

        st.caption(
            f"{subject} • {grade} • "
            f"{duration} minutes • "
            f"{difficulty} difficulty"
        )

        tab1, tab2, tab3, tab4 = st.tabs(
            [
                "📚 Lesson Plan",
                "🧠 Teaching Content",
                "📝 Assessment",
                "🔍 Review"
            ]
        )

        with tab1:

            st.markdown(
                package["lesson_plan"]
            )

        with tab2:

            st.markdown(
                package["content"]
            )

        with tab3:

            st.markdown(
                package["assessment"]
            )

        with tab4:

            st.markdown(
                package["review"]
            )


        # =================================================
        # DOWNLOAD PACKAGE
        # =================================================

        download_text = f"""
EDUAGENT AI — TEACHING PACKAGE

Subject: {subject}
Grade: {grade}
Topic: {topic}
Class Duration: {duration} minutes
Difficulty: {difficulty}

==================================================
LESSON PLAN
==================================================

{package["lesson_plan"]}

==================================================
TEACHING CONTENT
==================================================

{package["content"]}

==================================================
ASSESSMENT
==================================================

{package["assessment"]}

==================================================
REVIEW
==================================================

{package["review"]}
"""

        def create_teaching_package_pdf(
    subject,
    grade,
    topic,
    duration,
    difficulty,
    package
):
    buffer = BytesIO()

    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=45,
        leftMargin=45,
        topMargin=45,
        bottomMargin=45
    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "TitleStyle",
        parent=styles["Title"],
        alignment=TA_CENTER,
        fontSize=20,
        spaceAfter=15
    )

    heading_style = ParagraphStyle(
        "HeadingStyle",
        parent=styles["Heading2"],
        fontSize=15,
        spaceBefore=12,
        spaceAfter=8
    )

    body_style = ParagraphStyle(
        "BodyStyle",
        parent=styles["BodyText"],
        fontSize=10,
        leading=14,
        spaceAfter=6
    )

    story = []

    story.append(
        Paragraph(
            "EduAgent AI - Teaching Package",
            title_style
        )
    )

    story.append(
        Paragraph(
            f"<b>Subject:</b> {subject}<br/>"
            f"<b>Grade:</b> {grade}<br/>"
            f"<b>Topic:</b> {topic}<br/>"
            f"<b>Class Duration:</b> {duration} minutes<br/>"
            f"<b>Difficulty:</b> {difficulty}",
            body_style
        )
    )

    sections = [
        ("Lesson Plan", package["lesson_plan"]),
        ("Teaching Content", package["content"]),
        ("Assessment", package["assessment"]),
        ("Review", package["review"])
    ]

    for heading, content in sections:

        story.append(
            Paragraph(
                heading,
                heading_style
            )
        )

        safe_content = (
            str(content)
            .replace("&", "&amp;")
            .replace("<", "&lt;")
            .replace(">", "&gt;")
            .replace("\n", "<br/>")
        )

        story.append(
            Paragraph(
                safe_content,
                body_style
            )
        )

    doc.build(story)

    buffer.seek(0)

    return buffer.getvalue()


pdf_data = create_teaching_package_pdf(
    subject,
    grade,
    topic,
    duration,
    difficulty,
    package
)

st.download_button(
    "📥 Download Teaching Package as PDF",
    data=pdf_data,
    file_name=(
        f"EduAgent_"
        f"{topic.replace(' ', '_')}.pdf"
    ),
    mime="application/pdf",
    type="secondary",
    use_container_width=True
)


# =========================================================
# FOOTER
# =========================================================

st.markdown("---")

st.caption(
    "🎓 EduAgent AI | Multi-Agent Teaching Assistant | "
    "Multi-Format Knowledge Base | Simple RAG | "
    "Visual Understanding"
)
